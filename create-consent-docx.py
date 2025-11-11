from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Create a new Document
doc = Document()

# Title
title = doc.add_heading('DylianAI – SMS Consent Collection Policy', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Business Information
doc.add_paragraph('Business Name: DylianAI')
doc.add_paragraph('Business Address: 21329 N 73rd Way, Scottsdale, AZ 85255-4747')
doc.add_paragraph('Business Phone: (877) 446-4085')
doc.add_paragraph('Business Email: contact@dylianai.com')
doc.add_paragraph('')
doc.add_paragraph('Website: https://dylianai.com/')
doc.add_paragraph('')

# Overview Section
doc.add_heading('Overview', level=1)
doc.add_paragraph(
    'DylianAI uses SMS text messaging to provide appointment confirmations, reminders, and support updates to clients. '
    'All messages are strictly transactional in nature. DylianAI does not send promotional or marketing text messages.'
)
doc.add_paragraph('')

# Consent Collection Process Section
doc.add_heading('Consent Collection Process', level=1)
doc.add_paragraph('Clients initiate contact with DylianAI by phone to schedule or manage an appointment.')
doc.add_paragraph('')
doc.add_paragraph('During the call, our representative asks the client for verbal consent to receive SMS reminders or updates.')
doc.add_paragraph('')
doc.add_paragraph('Example consent script:')
doc.add_paragraph('')
consent_script = doc.add_paragraph(
    '"Before we send you text messages with appointment reminders or updates, do we have your permission to text you '
    'at this number? Message and data rates may apply, and you can reply STOP at any time to opt out."'
)
consent_script.italic = True
doc.add_paragraph('')
doc.add_paragraph(
    'Once the client verbally agrees, the representative records this consent in our scheduling or CRM system with a note such as:'
)
doc.add_paragraph('')
consent_note = doc.add_paragraph('"Verbal SMS Consent: Yes – obtained on [date]."')
consent_note.italic = True
doc.add_paragraph('')
doc.add_paragraph('The client is then added to our SMS system for appointment-related notifications only.')
doc.add_paragraph('')

# Opt-Out and Help Instructions Section
doc.add_heading('Opt-Out and Help Instructions', level=1)
doc.add_paragraph('Each outgoing message includes:')
doc.add_paragraph('')
doc.add_paragraph('Opt-out: "Reply STOP to unsubscribe."', style='List Bullet')
doc.add_paragraph('Help: "Reply HELP for support."', style='List Bullet')
doc.add_paragraph('')
doc.add_paragraph('Opt-out requests are processed automatically and honored immediately.')
doc.add_paragraph('')

# Message Frequency Section
doc.add_heading('Message Frequency', level=1)
doc.add_paragraph(
    'Messages are only sent when necessary for appointment confirmation, reminder, or rescheduling. '
    'Frequency depends on the number of appointments a client schedules.'
)
doc.add_paragraph('')

# Privacy Commitment Section
doc.add_heading('Privacy Commitment', level=1)
doc.add_paragraph(
    'Client contact information is stored securely and used only for communication related to scheduled appointments. '
    'DylianAI does not share or sell contact data to third parties.'
)
doc.add_paragraph('')

# Digital Signature Section
doc.add_heading('Digital Signature', level=1)
doc.add_paragraph('Signed electronically by DylianAI')
doc.add_paragraph('Date: ______________________')

# Save the document
doc.save('DylianAI_SMS_Consent_Collection_Policy.docx')
print('Word document created successfully: DylianAI_SMS_Consent_Collection_Policy.docx')

